"""Pure, deterministic presentation of already retained report snapshots."""

import base64
import binascii
import hashlib
import html
import json
import re
import unicodedata
from datetime import datetime
from io import BytesIO
from math import ceil
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._writer import WorksheetWriter
from openpyxl.writer.excel import ExcelWriter

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_MAX_BYTES = 16_000_000
_FIXED_TIME = datetime(2000, 1, 1)
_AUTHORITY = (
    "Retained analytical report. Current use authorized: false. Business effect authorized: false."
)
_JSON = dict(ensure_ascii=True, sort_keys=True, separators=(",", ":"), allow_nan=False)


class _MemoryWriter(ExcelWriter):
    """Keep openpyxl worksheet serialization in memory, including its XML staging."""

    def write_worksheet(self, ws):
        ws._drawing = SpreadsheetDrawing()
        writer = WorksheetWriter(ws, out=BytesIO())
        writer.write()
        ws._rels = writer._rels
        self._archive.writestr(ws.path[1:], writer.read())
        self.manifest.append(ws)


def _json(value):
    return json.dumps(value, **_JSON)


def _scalar(value):
    if value is None:
        return "[Null value]"
    if value == "":
        return "[Empty text]"
    if isinstance(value, bool):
        return "True" if value else "False"
    return str(value)


def _display(value, field):
    state = value.get("state", "VALUE")
    if state == "MISSING":
        return "[Missing]"
    if state == "NULL":
        return "[Null]"
    raw = value.get("value")
    if raw is not None and raw != "" and field["kind"] == "reference" and value.get("label"):
        return value["label"]
    return _scalar(raw)


def _fields(section):
    fields = {item["key"]: item for item in section["projection"]["descriptor"]["fields"]}
    return [fields[key] for key in section["reference"]["columns"]]


def _context(section):
    descriptor = section["projection"]["descriptor"]
    context = [(item["label"], item["value"]) for item in descriptor.get("context", [])]
    supplied = {label.casefold().strip() for label, _ in context}
    context.extend(
        (label, "Not supplied")
        for label in ("Period", "Book", "Currency")
        if label.casefold() not in supplied
    )
    context.extend(
        (label, descriptor[key])
        for label, key in (
            ("Valid at", "valid_at"),
            ("Known at", "known_at"),
            ("Recorded at", "recorded_at"),
        )
    )
    context.extend(("Coverage: " + item["label"], item["value"]) for item in descriptor["coverage"])
    return context


def _advanced(snapshot):
    """Select provenance explicitly; never serialize hidden values or evidence cells."""
    yield (
        "Report",
        "",
        "",
        {
            key: snapshot[key]
            for key in (
                "contract",
                "company",
                "company_label",
                "current_use_authorized",
                "business_effect_authorized",
            )
        },
    )
    yield (
        "Report",
        "",
        "Composition",
        {
            key: snapshot["composition"].get(key)
            for key in ("company_id", "valid_at", "known_at", "title", "commentary")
        },
    )
    for index, section in enumerate(snapshot["sections"], 1):
        name = str(index)
        descriptor = section["projection"]["descriptor"]
        reference = section["reference"]
        # Filter values are omitted when their field is not selected for export.
        selected_reference = dict(
            reference,
            filters=[
                item
                for item in reference.get("filters", [])
                if item["field"] in reference["columns"]
            ],
        )
        yield name, "", "Section reference", selected_reference
        yield (
            name,
            "",
            "Retained authority",
            {
                key: descriptor.get(key)
                for key in (
                    "function",
                    "company",
                    "run_id",
                    "authority",
                    "definitions",
                    "unavailable_operations",
                    "current_use_authorized",
                    "business_effect_authorized",
                )
            },
        )
        yield name, "", "Authority observation", section["authority_observation"]
        for field in _fields(section):
            yield (
                name,
                "",
                "Field: " + field["key"],
                {key: value for key, value in field.items() if key != "options"},
            )
        for row in section["projection"]["rows"]:
            yield (
                name,
                row["key"],
                "Row reference",
                {
                    "trace": row["trace"],
                    "contributor_count": row["contributor_count"],
                },
            )
            for field in _fields(section):
                key = field["key"]
                value = row["values"].get(key, {"state": "MISSING"})
                raw = value.get("value")
                yield (
                    name,
                    row["key"],
                    key,
                    {
                        "field_kind": field["kind"],
                        "value_key_present": "value" in value,
                        "value_type": "null" if raw is None else type(raw).__name__,
                        "selected_cell_present": key in row["values"],
                        "cell": value,
                    },
                )
            for contributor in section.get("contributors", {}).get(row["key"], []):
                yield (
                    name,
                    row["key"],
                    "Source reference",
                    {
                        key: contributor[key]
                        for key in (
                            "reference",
                            "document_id",
                            "source_sha256",
                            "sheet",
                            "coordinate",
                            "basis",
                        )
                        if key in contributor
                    },
                )


def _excel_text(value):
    # XML 1.0 cannot represent these controls; show their exact escape spelling.
    value = re.sub(
        r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]", lambda m: f"\\u{ord(m[0]):04x}", str(value)
    )
    if len(value) > 32767:
        return value[:32720] + " [continued in Advanced]"
    return value


def _append(sheet, values, *, heading=False):
    sheet.append([_excel_text(value) for value in values])
    for column in range(1, len(values) + 1):
        cell = sheet.cell(sheet._current_row, column)
        cell.data_type = "s"  # Includes =, +, -, @, external links and source formulas.
        cell.number_format = "@"
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = Font(
            name="Calibri", size=11, color="FFFFFF" if heading else "172F3E", bold=heading
        )
        if heading:
            cell.fill = PatternFill("solid", fgColor="173B49")


def _style(sheet, width=28, widths=None):
    sheet.sheet_view.showGridLines = False
    for column in range(1, sheet.max_column + 1):
        letter = get_column_letter(column)
        sheet.column_dimensions[letter].width = (widths or {}).get(letter, width)
    for row in sheet:
        lines = 1
        for cell in row:
            capacity = max(1, sheet.column_dimensions[cell.column_letter].width - 3)
            wrapped = sum(
                max(
                    1,
                    ceil(
                        sum(2 if unicodedata.east_asian_width(c) in {"F", "W"} else 1 for c in line)
                        / capacity
                    ),
                )
                for line in str(cell.value or "").split("\n")
            )
            lines = max(lines, wrapped)
        sheet.row_dimensions[row[0].row].height = min(409, 16 * lines + 6)
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.print_area = sheet.dimensions


def _xlsx(snapshot):
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Report"
    workbook.properties.creator = "G8"
    workbook.properties.lastModifiedBy = "G8"
    workbook.properties.created = _FIXED_TIME
    workbook.properties.modified = _FIXED_TIME
    composition = snapshot["composition"]
    _append(overview, [composition["title"]], heading=True)
    for label, value in (
        ("Company", snapshot["company_label"]),
        ("Company valid at", composition["valid_at"]),
        ("Company known at", composition["known_at"]),
        ("Author commentary", composition.get("commentary", "")),
        ("Authority", _AUTHORITY),
        (
            "Values",
            "Exact retained text. [Missing], [Null], [Null value], and [Empty text] are distinct. "
            "Advanced preserves original types and states.",
        ),
    ):
        _append(overview, [label, value])
    _append(overview, ["Section", "Title"], heading=True)
    for index, section in enumerate(snapshot["sections"], 1):
        _append(overview, [str(index), section["reference"]["title"]])
        sheet = workbook.create_sheet(f"Section {index}")
        _append(sheet, [section["reference"]["title"]], heading=True)
        _append(sheet, ["Company", snapshot["company_label"]])
        for label, value in _context(section):
            _append(sheet, [label, value])
        fields = _fields(section)
        _append(
            sheet,
            [
                field["label"] + (f" ({field['unit']})" if field.get("unit") else "")
                for field in fields
            ],
            heading=True,
        )
        header = sheet.max_row
        for row in section["projection"]["rows"]:
            _append(
                sheet,
                [
                    _display(row["values"].get(field["key"], {"state": "MISSING"}), field)
                    for field in fields
                ],
            )
        sheet.freeze_panes = f"A{header + 1}"
        sheet.print_title_rows = f"{header}:{header}"
        sheet.auto_filter.ref = f"A{header}:{get_column_letter(len(fields))}{sheet.max_row}"
        _style(sheet)
    advanced = workbook.create_sheet("Advanced")
    _append(
        advanced,
        ["Section", "Row key", "Record", "Part", "Exact JSON (join parts in order)"],
        heading=True,
    )
    for section, row_key, record, value in _advanced(snapshot):
        exact = _json(value)
        for start in range(0, len(exact), 30000):
            _append(
                advanced,
                [section, row_key, record, str(start // 30000 + 1), exact[start : start + 30000]],
            )
    advanced.freeze_panes = "A2"
    advanced.print_title_rows = "1:1"
    _style(advanced, widths={"E": 100})
    _style(overview, 36, widths={"B": 100})
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        _MemoryWriter(workbook, archive).write_data()
    result = BytesIO()
    with (
        ZipFile(BytesIO(stream.getvalue())) as source,
        ZipFile(result, "w", ZIP_DEFLATED, compresslevel=9) as target,
    ):
        for name in sorted(source.namelist()):
            content = source.read(name)
            if name == "docProps/core.xml":
                content = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>2000-01-01T00:00:00Z\2",
                    content,
                )
            member = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            member.compress_type = ZIP_DEFLATED
            member.create_system = 3
            member.external_attr = 0o600 << 16
            target.writestr(member, content)
    return result.getvalue()


def _html(snapshot):
    def escape(value):
        return html.escape(str(value), quote=True)

    composition = snapshot["composition"]
    parts = [
        '<!doctype html><html lang="en"><head><meta charset="utf-8">',
        '<meta name="viewport" content="width=device-width,initial-scale=1">',
        '<meta http-equiv="Content-Security-Policy" '
        "content=\"default-src 'none'; style-src 'unsafe-inline'; "
        "base-uri 'none'; form-action 'none'\">",
        "<title>" + escape(composition["title"]) + "</title>",
        "<style>body{font:15px/1.5 system-ui,sans-serif;color:#172f3e;max-width:1200px;"
        "margin:40px auto;padding:0 24px}h1,h2{line-height:1.2;color:#173b49}"
        "h1{font-size:32px}h2{margin-top:36px}dl{display:grid;"
        "grid-template-columns:minmax(130px,1fr) 4fr;gap:4px 18px}dt{font-weight:600}dd{margin:0}"
        "table{border-collapse:collapse;width:100%;margin:18px 0;table-layout:fixed}"
        "th,td{padding:9px 12px;border-bottom:1px solid #d9e1e5;text-align:left;"
        "vertical-align:top;overflow-wrap:anywhere;white-space:pre-wrap}"
        "th{background:#173b49;color:white}tbody tr:nth-child(even){background:#f3f6f7}"
        "p,dd,pre{overflow-wrap:anywhere;white-space:pre-wrap}pre{font-size:11px}"
        ".note{font-size:12px;color:#506572}thead{display:table-header-group}"
        "@media print{@page{size:A4 landscape;margin:14mm}"
        "body{margin:0;padding:0;max-width:none;font-size:10pt}h2{break-after:avoid}"
        "tr{break-inside:avoid}th{print-color-adjust:exact;-webkit-print-color-adjust:exact}"
        ".advanced{break-before:page}table{table-layout:auto}}</style></head><body>",
        "<h1>" + escape(composition["title"]) + "</h1>",
        "<p>" + escape(snapshot["company_label"]) + "</p>",
        "<dl><dt>Company valid at</dt><dd>"
        + escape(composition["valid_at"])
        + "</dd><dt>Company known at</dt><dd>"
        + escape(composition["known_at"])
        + "</dd></dl>",
        "<h2>Author commentary</h2><p>" + escape(composition.get("commentary", "")) + "</p>",
        '<p class="note">' + _AUTHORITY + "</p>",
        '<p class="note">Exact retained values. [Missing], [Null], [Null value], '
        "and [Empty text] are distinct. Advanced preserves original types and states.</p>",
    ]
    for section in snapshot["sections"]:
        parts.append("<section><h2>" + escape(section["reference"]["title"]) + "</h2><dl>")
        for label, value in _context(section):
            parts.append("<dt>" + escape(label) + "</dt><dd>" + escape(value) + "</dd>")
        parts.append("</dl><table><thead><tr>")
        fields = _fields(section)
        for field in fields:
            label = field["label"] + (f" ({field['unit']})" if field.get("unit") else "")
            parts.append('<th scope="col">' + escape(label) + "</th>")
        parts.append("</tr></thead><tbody>")
        for row in section["projection"]["rows"]:
            parts.append("<tr>")
            for field in fields:
                value = row["values"].get(field["key"], {"state": "MISSING"})
                parts.append("<td>" + escape(_display(value, field)) + "</td>")
            parts.append("</tr>")
        parts.append("</tbody></table></section>")
    parts.append(
        '<section class="advanced"><h2>Advanced</h2>'
        "<p>Exact selected values, types, states and source references.</p>"
    )
    for section, row_key, record, value in _advanced(snapshot):
        parts.append(
            '<p class="note">'
            + escape(f"{section} / {row_key} / {record}")
            + "</p><pre>"
            + escape(_json(value))
            + "</pre>"
        )
    parts.append("</section></body></html>")
    return "".join(parts).encode("utf-8")


def validate_artifact_bytes(artifact: dict) -> bytes:
    """Decode saved bytes, rejecting corrupt length, digest or base64."""
    try:
        content = base64.b64decode(artifact["content_base64"], validate=True)
    except (KeyError, ValueError, TypeError, binascii.Error) as exc:
        raise ValueError("Invalid retained report artifact encoding") from exc
    if not 0 < len(content) <= _MAX_BYTES or artifact.get("size_bytes") != len(content):
        raise ValueError("Invalid retained report artifact size")
    if artifact.get("sha256") != hashlib.sha256(content).hexdigest():
        raise ValueError("Invalid retained report artifact digest")
    return content


def build(snapshot: dict) -> dict:
    """Render a validated frozen snapshot without queries or new calculations."""
    artifacts = {}
    for kind, media_type, content in (
        ("xlsx", _XLSX, _xlsx(snapshot)),
        ("html", "text/html; charset=utf-8", _html(snapshot)),
    ):
        if not 0 < len(content) <= _MAX_BYTES:
            raise ValueError("Retained report export exceeds the artifact size limit")
        artifacts[kind] = {
            "media_type": media_type,
            "filename": f"retained-report.{kind}",
            "size_bytes": len(content),
            "sha256": hashlib.sha256(content).hexdigest(),
            "content_base64": base64.b64encode(content).decode("ascii"),
        }
    return artifacts
