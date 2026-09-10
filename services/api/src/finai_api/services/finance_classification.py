"""Evidence-backed rule execution returns proposals, never authoritative journal entries."""

import csv
import json
import re
from datetime import date, datetime
from decimal import Decimal
from hashlib import sha256
from io import BytesIO, StringIO
from typing import Any
from uuid import UUID

from psycopg.rows import dict_row

from finai_api.domain.finance_classification import (
    ClassificationRequest,
    ClassificationRule,
    FinanceClassificationPolicy,
)
from finai_api.domain.resource_lifecycle import VersionReference
from finai_api.domain.resources import ResourceMutation
from finai_api.domain.review import Principal
from finai_api.security import require_permission
from finai_api.services import fact_runs, ontology_definitions
from finai_api.services.certification import _current
from finai_api.services.resources import resource_connection
from finai_api.services.source_documents import document_bytes
from finai_api.services.upstream_authority import upstream_authority
from finai_api.services.workspace import WorkspaceError


def validate_policy(item: ResourceMutation, target: Any) -> None:
    policy = FinanceClassificationPolicy.model_validate(item.attributes["definition"])
    owner = str(item.resource_id)
    pins = {"company": policy.company, "chart": policy.chart, "evidence": policy.evidence}
    pins.update({"account:" + rule.key: rule.account for rule in policy.rules})
    loaded = {
        name: target(
            str(pin.resource_id), owner, "FINANCE_CLASSIFICATION:" + name, str(pin.version_id)
        )
        for name, pin in pins.items()
    }
    for name, kind in (
        ("company", "LegalEntity"),
        ("chart", "LocalChartOfAccounts"),
        ("evidence", "SourceEvidence"),
    ):
        if loaded[name]["object_type"] != kind:
            raise WorkspaceError(422, f"Classification {name} must reference {kind}")
    if str(loaded["chart"]["attributes"].get("legal_entity_id")) != str(policy.company.resource_id):
        raise WorkspaceError(422, "Classification chart belongs to another company")
    for rule in policy.rules:
        account = loaded["account:" + rule.key]
        if account["object_type"] != "LocalAccount" or str(
            account["attributes"].get("chart_id")
        ) != str(policy.chart.resource_id):
            raise WorkspaceError(422, "Classification account belongs to another chart")
    if item.access_entity is not None and item.access_entity != policy.exact_scope.legal_entity_id:
        raise WorkspaceError(
            422, "Classification policy access scope differs from its company scope"
        )


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value)


def source_rows(
    content: bytes, policy: FinanceClassificationPolicy, offset: int, limit: int
) -> tuple[list[dict], bool]:
    if len(content) > 32_000_000:
        raise WorkspaceError(422, "Classification source exceeds its bounded size")
    rows: list[list[Any]]
    if policy.source_format == "csv":
        matrix = csv.reader(StringIO(content.decode("utf-8-sig"), newline=""))
        rows = []
        for index, row in enumerate(matrix, 1):
            if index > 30100:
                raise WorkspaceError(422, "Source exceeds 30000 row classification bound")
            rows.append(list(row))
    elif policy.source_format == "json":
        value = json.loads(content)
        if (
            not isinstance(value, list)
            or len(value) > 30000
            or any(not isinstance(row, dict) for row in value)
        ):
            raise WorkspaceError(422, "JSON classification source must be an array of row objects")
        if any(set(row) != set(policy.columns) for row in value):
            raise WorkspaceError(
                422, "Retained JSON columns differ from the reviewed source policy"
            )
        return (
            [
                {
                    "source_row": index + 1,
                    "values": {key: _text(row[key]) for key in policy.columns},
                }
                for index, row in enumerate(value)
                if offset <= index < offset + limit
            ],
            len(value) > offset + limit,
        )
    elif policy.source_format == "xlsx":
        import openpyxl

        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False)
        try:
            if policy.sheet not in workbook.sheetnames:
                raise WorkspaceError(422, "Reviewed worksheet is absent from the retained source")
            sheet = workbook[policy.sheet]
            if sheet.max_row > 30100 or sheet.max_column > 256:
                raise WorkspaceError(422, "Workbook exceeds the classification bound")
            rows = [
                [None if cell.data_type == "f" else _text(cell.value) for cell in row]
                for row in sheet.iter_rows()
            ]
        finally:
            workbook.close()
    else:
        import xlrd

        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
        try:
            if policy.sheet not in workbook.sheet_names():
                raise WorkspaceError(422, "Reviewed worksheet is absent from the retained source")
            sheet = workbook.sheet_by_name(policy.sheet)
            if sheet.nrows > 30100 or sheet.ncols > 256:
                raise WorkspaceError(422, "Workbook exceeds the classification bound")
            rows = [
                [_text(cell) for cell in sheet.row_values(index)] for index in range(sheet.nrows)
            ]
        finally:
            workbook.release_resources()
    if len(rows) < policy.header_row or rows[policy.header_row - 1] != policy.columns:
        raise WorkspaceError(422, "Retained source header differs from the reviewed policy")
    body = rows[policy.header_row :]
    return (
        [
            {
                "source_row": policy.header_row + index + 1,
                "values": dict(zip(policy.columns, row, strict=True))
                if len(row) == len(policy.columns)
                else {},
                "width_valid": len(row) == len(policy.columns),
            }
            for index, row in enumerate(body)
            if offset <= index < offset + limit
        ],
        len(body) > offset + limit,
    )


def classify_values(policy: FinanceClassificationPolicy, values: dict[str, Any]) -> dict[str, Any]:
    matches = []
    for candidate_rule in policy.rules:
        applicable = True
        for condition in candidate_rule.conditions:
            value = values.get(condition.field)
            if not isinstance(value, str):
                applicable = False
                break
            expected = condition.value
            if not condition.case_sensitive:
                value, expected = value.casefold(), expected.casefold()
            applicable &= (
                value == expected
                if condition.operator == "equals"
                else expected in value
                if condition.operator == "contains"
                else value.startswith(expected)
            )
        if applicable:
            matches.append(candidate_rule)
    missing = []
    amount = values.get(policy.amount_field)
    pattern = r"-?\d+(?:\.\d+)?" if policy.decimal_format == "DOT" else r"-?\d+(?:,\d+)?"
    if not isinstance(amount, str) or len(amount) > 50 or re.fullmatch(pattern, amount) is None:
        missing.append("amount")
        amount = None
    else:
        amount = str(Decimal(amount.replace(",", ".")))
    period_date = values.get(policy.date_field)
    try:
        if not isinstance(period_date, str):
            raise ValueError("Effective date is not text")
        parsed_date = date.fromisoformat(period_date)
        if parsed_date.strftime("%Y-%m") != policy.exact_scope.period:
            missing.append("date_outside_scope")
    except (TypeError, ValueError):
        missing.append("effective_date")
    if values.get(policy.currency_field) != policy.exact_scope.currency:
        missing.append("currency")
    if not values.get(policy.identity_field):
        missing.append("source_identity")
    dimensions = {
        name: values.get(field) or None for name, field in policy.dimension_fields.items()
    }
    missing.extend(
        "dimension:" + name for name in policy.required_dimensions if dimensions.get(name) is None
    )
    rule: ClassificationRule | None = matches[0] if len(matches) == 1 else None
    if not matches:
        missing.append("NO_MATCHING_RULE")
    elif len(matches) > 1:
        missing.append("AMBIGUOUS_RULES")
    return {
        "state": "CANDIDATE" if rule and not missing else "UNAVAILABLE",
        "epistemic": "INFERRED",
        "matching_rules": [row.key for row in matches],
        "missing": missing,
        "account": rule.account.model_dump(mode="json") if rule else None,
        "grain": rule.grain if rule else None,
        "side": rule.side if rule else None,
        "process_type": rule.process_type if rule else None,
        "amount": amount,
        "currency": values.get(policy.currency_field),
        "effective_date": period_date,
        "dimensions": dimensions,
        "dimension_authority": "OBSERVED_LABELS_ONLY",
        "reason": rule.rationale if rule else "No unique evidenced interpretation",
        "accounting_use_authorized": False,
        "canonical_posting_created": False,
    }


def classify(principal: Principal, request: ClassificationRequest) -> dict[str, Any]:
    require_permission(principal, "ontology_read")
    row = ontology_definitions.definition(
        principal, request.policy.resource_id, request.policy.version_id
    )
    if row["object_type"] != "FinanceClassificationPolicy":
        raise WorkspaceError(422, "Classification requires an exact reviewed policy")
    policy = FinanceClassificationPolicy.model_validate(row["attributes"]["definition"])
    if (
        policy.exact_scope != principal.scope
        or row["access_entity"] != principal.scope.legal_entity_id
    ):
        raise WorkspaceError(409, "Classification policy belongs to another exact scope")
    pins = {entry["relation"]: entry for entry in row["dependencies"]}
    with resource_connection(principal) as conn, conn.cursor(row_factory=dict_row) as cursor:
        _current(cursor, principal, request.policy)
        upstream_authority(cursor, principal.scope.tenant_id, request.policy.version_id)

        def target(identity: str, owner: str, relation: str, version: str) -> dict:
            pin = pins.get(relation)
            if not pin or str(pin["resource_id"]) != identity or str(pin["version_id"]) != version:
                raise WorkspaceError(409, "Classification policy dependency changed")
            return _current(
                cursor,
                principal,
                VersionReference(resource_id=UUID(identity), version_id=UUID(version)),
            )

        validate_policy(
            ResourceMutation(
                resource_id=request.policy.resource_id,
                object_type=row["object_type"],
                identity_key=row["identity_key"],
                display_name=row["display_name"],
                attributes=row["attributes"],
                valid_from=row["valid_from"],
                access_entity=row["access_entity"],
            ),
            target,
        )
    document, content = document_bytes(principal, request.document_id)
    try:
        rows, more = source_rows(content, policy, request.offset, request.limit)
    except (ValueError, UnicodeError, KeyError) as exc:
        raise WorkspaceError(
            422, "Retained source cannot be read using the reviewed layout"
        ) from exc
    output = [
        {**source, "classification": classify_values(policy, source["values"])} for source in rows
    ]
    return fact_runs.retain_run(
        principal,
        {
            "operation": "classify_fact",
            "state": "CANDIDATE_ONLY",
            "policy": request.policy.model_dump(mode="json"),
            "source": {
                "document_id": request.document_id,
                "sha256": sha256(content).hexdigest(),
                "filename": document.get("filename"),
            },
            "source_family": policy.source_family,
            "sheet": policy.sheet,
            "rows": output,
            "offset": request.offset,
            "next_offset": request.offset + request.limit if more else None,
            "accounting_use_authorized": False,
        },
        runtime="finance-classification/1",
    )
