# ruff: noqa: F811
"""HTTP-to-retention and deterministic workflow failure contracts with controlled transport."""

import asyncio
import base64
import json
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from temporalio.exceptions import ActivityError, WorkflowAlreadyStartedError
from test_definition_history import DB, item, retained  # noqa: F401
from test_function_execution import function_case

from finai_api import transformation_workflow as orchestration
from finai_api.api import transformation_routes as builds
from finai_api.api import workflow_routes as reports
from finai_api.config import get_settings
from finai_api.domain.transformation import TransformationRunRequest
from finai_api.main import app
from finai_api.services import report_workflows as records
from finai_api.services import transformation_runs as runs
from finai_api.services.workspace import WorkspaceError


class Transport:
    def __init__(self):
        self.starts = []
        self.signals = []
        self.runtime_status = "RUNNING"
        self.state = "WAITING_REVIEW"

    async def start_workflow(self, function, context, **options):
        self.starts.append((context, options))
        if len(self.starts) > 1:
            raise WorkflowAlreadyStartedError(options["id"], "fixture-existing")

    def get_workflow_handle(self, identity):
        self.identity = identity
        return self

    async def describe(self):
        return SimpleNamespace(
            status=(SimpleNamespace(name=self.runtime_status) if self.runtime_status else None)
        )

    async def query(self, query):
        return {"state": self.state, "result": {}}

    async def signal(self, name, *args):
        self.signals.append((self.identity, name, args))


async def offline():
    raise WorkspaceError(503, "Controlled fixture runtime unavailable")


def http(monkeypatch, actor):
    checker = actor.model_copy(update={"actor_id": "independent-api-fixture-checker"})
    denied = actor.model_copy(update={"permissions": ("read", "ontology_read")})
    monkeypatch.setenv(
        "FINAI_ACCESS_TOKENS",
        json.dumps(
            {
                "fixture-maker": actor.model_dump(mode="json"),
                "fixture-checker": checker.model_dump(mode="json"),
                "fixture-read-only": denied.model_dump(mode="json"),
            }
        ),
    )
    get_settings.cache_clear()
    return TestClient(app, headers={"Authorization": "Bearer fixture-maker"})


def transport(monkeypatch):
    connection = Transport()

    async def connect():
        return connection

    monkeypatch.setattr(reports, "client", connect)
    monkeypatch.setattr(builds, "client", connect)
    return connection


@DB
def test_report_http_retains_before_start_and_controls_require_independent_review(
    retained, monkeypatch
):
    actor = retained[0].model_copy(update={"permissions": ("read", "ingest", "review", "export")})
    api = http(monkeypatch, actor)
    source = api.post(
        "/v1/hydration/ingest",
        json={
            "scope": actor.scope.model_dump(mode="json"),
            "filename": "synthetic-workflow.csv",
            "csv_text": "label,value\nfixture,unavailable\n",
        },
    )
    assert source.status_code == 200, source.text
    request = {
        "report": {
            "period": "2026-08",
            "company_label": "Synthetic report fixture",
            "currency": "GEL",
            "receipt_ids": [source.json()["receipt_id"]],
        }
    }
    monkeypatch.setattr(reports, "client", offline)
    assert api.post("/v1/workspace/workflows", json=request).status_code == 503
    retained_list = api.get("/v1/workspace/workflows")
    identity = retained_list.json()[0]["workflow_id"]
    assert api.get(f"/v1/workspace/workflows/{identity}").json()["runtime_status"] == "UNOBSERVABLE"
    runtime = transport(monkeypatch)
    for _ in range(2):
        assert api.post("/v1/workspace/workflows", json=request).json() == {"workflow_id": identity}
    assert runtime.starts[0][0]["scope"] == actor.scope.model_dump(mode="json")
    assert "report" not in runtime.starts[0][0]  # Source bytes stay in G8, not Temporal history.
    status = api.get(f"/v1/workspace/workflows/{identity}").json()
    assert status["runtime_status"] == "RUNNING" and status["publications"] == []
    runtime.runtime_status = None
    assert api.get(f"/v1/workspace/workflows/{identity}").json()["runtime_status"] == "UNKNOWN"
    endpoint = f"/v1/workspace/workflows/{identity}/control"
    command = {
        "command": "complete",
        "reason": "Review the synthetic source assessment",
        "idempotency_key": str(uuid4()),
    }
    assert api.post(endpoint, json=command).status_code == 403
    checker = {"Authorization": "Bearer fixture-checker"}
    runtime.state = "FAILED"
    assert api.post(endpoint, json=command, headers=checker).status_code == 409
    runtime.state = "WAITING_REVIEW"
    assert api.post(endpoint, json=command, headers=checker).status_code == 200
    assert api.post(endpoint, json=command, headers=checker).status_code == 200
    assert (
        api.post(endpoint, json={**command, "command": "cancel"}, headers=checker).status_code
        == 409
    )
    for name in ("pause", "resume", "retry", "cancel"):
        assert (
            api.post(
                endpoint, json={**command, "command": name, "idempotency_key": str(uuid4())}
            ).status_code
            == 200
        )
    assert len(records.read(actor, identity)["events"]) == 5
    assert (
        api.post(
            endpoint,
            json={**command, "command": "pause"},
            headers={"Authorization": "Bearer fixture-read-only"},
        ).status_code
        == 403
    )


@DB
def test_transformation_http_control_and_retained_review_survive_offline_transport(
    retained, monkeypatch
):
    reader, invocation, _, _, _ = function_case(retained)
    actor = reader.model_copy(
        update={"permissions": (*reader.permissions, "read", "ingest", "review")}
    )
    definition = item(
        "TransformationDefinition",
        {
            "resource_budget": {
                "max_returned_rows": 10,
                "max_derived_evaluations": 0,
                "max_published_result_bytes": 1000000,
            },
            "publication_review": {"question": "Review synthetic observation publication only"},
            "definition": {
                "nodes": [
                    {
                        "node_id": "source",
                        "function_id": str(invocation.function.resource_id),
                        "limit": 10,
                    }
                ],
                "outputs": [{"output_id": "observations", "node_id": "source"}],
            },
        },
    )
    row = retained[1](definition)[0]
    request = TransformationRunRequest(
        transformation={"resource_id": row["resource_id"], "version_id": row["version_id"]},
        valid_at=invocation.valid_at,
        known_at=invocation.known_at,
    )
    api, runtime = http(monkeypatch, actor), transport(monkeypatch)
    base = "/v1/ontology/transformations"
    started = api.post(base + "/runs", json=request.model_dump(mode="json"))
    assert started.status_code == 200, started.text
    identity = started.json()["workflow_id"]
    endpoint = base + f"/runs/{request.request_id}"
    assert api.get(base).status_code == 200
    assert api.get(base + "/runs").status_code == 200
    assert api.get(endpoint).json()["runtime_status"] == "RUNNING"
    runtime.runtime_status = None
    assert api.get(endpoint).json()["runtime_status"] == "UNKNOWN"
    control = {
        "command": "pause",
        "reason": "Inspect the synthetic retained build",
        "idempotency_key": str(uuid4()),
    }
    assert api.post(endpoint + "/control", json=control).status_code == 409
    runtime.runtime_status = "RUNNING"
    for command in ("pause", "resume", "cancel"):
        payload = {**control, "command": command, "idempotency_key": str(uuid4())}
        # Completed runtime handles refuse cancellation; keep this review fixture active.
        if command == "cancel":
            runtime.runtime_status = "COMPLETED"
            assert api.post(endpoint + "/control", json=payload).status_code == 409
            runtime.runtime_status = "RUNNING"
            continue
        assert api.post(endpoint + "/control", json=payload).status_code == 200
        assert api.post(endpoint + "/control", json=payload).status_code == 200
        assert (
            api.post(endpoint + "/control", json={**payload, "command": "cancel"}).status_code
            == 409
        )
    context = {
        "workflow_id": identity,
        "actor_id": actor.actor_id,
        "scope": actor.scope.model_dump(mode="json"),
    }
    monkeypatch.setattr(records, "current_principal", lambda *_: actor)
    runs.execute_node({**context, "node_id": "source"})
    runs.publication_review(context)
    decision = {
        "decision_id": str(uuid4()),
        "decision": "APPROVED",
        "reason": "Independent review of synthetic observations",
    }
    assert api.post(endpoint + "/publication-review", json=decision).status_code == 403
    monkeypatch.setattr(builds, "client", offline)
    accepted = api.post(
        endpoint + "/publication-review",
        json=decision,
        headers={"Authorization": "Bearer fixture-checker"},
    )
    assert accepted.status_code == 200, accepted.text
    assert accepted.json()["runtime_notified"] is False
    assert accepted.json()["publication_review"]["state"] == "APPROVED"
    published = runs.publish(context)
    saved = api.get(endpoint).json()
    assert saved["runtime_status"] == "UNOBSERVABLE"
    assert saved["publications"][0]["publication_id"] == published["publication_id"]
    transport(monkeypatch)
    replay = api.post(
        endpoint + "/publication-review",
        json=decision,
        headers={"Authorization": "Bearer fixture-checker"},
    )
    assert replay.json()["runtime_notified"] is True
    assert api.get(f"/v1/workspace/workflows/{identity}").status_code == 409


@pytest.mark.parametrize(
    "model,payload",
    [
        (builds.Control, {"command": "pause", "idempotency_key": uuid4()}),
        (builds.ReviewDecision, {"decision": "APPROVED", "decision_id": uuid4()}),
    ],
)
def test_transformation_reasons_cannot_be_padding(model, payload):
    with pytest.raises(ValueError, match="non-padding"):
        model(**payload, reason=" " * 20)
    assert (
        model(**payload, reason="  Meaningful fixture reason  ").reason
        == "Meaningful fixture reason"
    )


def test_temporal_connection_failure_is_actionable_without_leaking_details(monkeypatch):
    async def fail(*args, **kwargs):
        raise RuntimeError("private transport details")

    monkeypatch.setattr(reports.Client, "connect", fail)
    with pytest.raises(WorkspaceError) as failure:
        asyncio.run(reports.client())
    assert failure.value.status == 503
    assert "retained requests can be retried" in failure.value.detail
    assert "private" not in failure.value.detail


@DB
def test_report_calculation_http_retains_real_workbook_and_exports_reference_only(
    retained, monkeypatch
):
    actor = retained[0].model_copy(update={"permissions": ("read", "ingest", "export")})
    api = http(monkeypatch, actor)
    book = Workbook()
    revenue = book.active
    revenue.title = "Revenue Breakdown"
    revenue.append(["Product", "Gross", "Discount", "Net Revenue"])
    revenue.append(["Euro Regular", 10, 2, "=B2-C2"])
    cost = book.create_sheet("COGS Breakdown")
    for cell, value in {
        "K1": "6",
        "L1": "7310",
        "O1": "8230",
        "A2": "Euro Regular",
        "K2": 3,
        "L2": 1,
        "O2": 0,
    }.items():
        cost[cell] = value
    output = BytesIO()
    book.save(output)
    source = api.post(
        "/v1/hydration/ingest",
        json={
            "scope": actor.scope.model_dump(mode="json"),
            "filename": "synthetic-reference.xlsx",
            "source_use": "HISTORICAL_REFERENCE",
            "xlsx_base64": base64.b64encode(output.getvalue()).decode(),
        },
    )
    assert source.status_code == 200, source.text
    request = {"receipt_id": source.json()["receipt_id"]}
    endpoint = "/v1/workspace/report-calculations"
    result = api.post(endpoint, json=request)
    assert result.status_code == 200, result.text
    original = result.json()
    assert original["state"] == "REFERENCE_ONLY"
    values = {row["id"]: row["amount"] for row in original["metrics"]}
    assert values["gross_profit"] == "4" and values["legacy_ebitda"] is None
    assert api.post(endpoint, json=request).json() == original
    path = endpoint + "/" + original["calculation_id"]
    assert api.get(path).json() == original
    exported = api.get(path + "/export")
    assert exported.status_code == 200 and exported.headers["cache-control"] == "no-store"
    rendered = load_workbook(BytesIO(exported.content))
    assert "NOT APPROVED" in rendered["Operating P&L"]["A1"].value
    rendered.close()
    assert api.get(path, headers={"Authorization": "Bearer fixture-read-only"}).status_code == 403
    assert api.get(endpoint + "/missing").status_code == 404
    malformed = api.post(
        "/v1/hydration/ingest",
        json={
            "scope": actor.scope.model_dump(mode="json"),
            "filename": "synthetic-unsupported.csv",
            "csv_text": "label,value\nfixture,unavailable\n",
        },
    )
    assert malformed.status_code == 200
    refused = api.post(endpoint, json={"receipt_id": malformed.json()["receipt_id"]})
    assert refused.status_code == 422


@pytest.mark.parametrize("outcome", ["approved", "rejected", "cancelled", "timeout"])
def test_transformation_publication_review_arbitrates_notifications_timeout_and_cancel(
    monkeypatch, outcome
):
    instance = orchestration.TransformationWorkflow()
    calls = []
    decisions = iter(["PENDING", "REJECTED" if outcome == "rejected" else "APPROVED"])

    async def execute(name, context, **options):
        calls.append(name)
        assert options["retry_policy"].maximum_attempts == 3
        if name == "transformation_load":
            return {"node_order": [], "dependencies": {}, "publication_review": True}
        if name == "transformation_publication_review":
            return {"state": next(decisions)}
        return {"publication_id": "fixture-reference-only"}

    async def wait(condition, **options):
        assert instance.state == "AWAITING_REVIEW"
        if outcome == "cancelled":
            instance.control({"id": "cancel", "command": "cancel"})
        elif outcome == "timeout":
            raise TimeoutError
        else:
            instance.review_changed()
        assert condition()

    monkeypatch.setattr(orchestration.workflow, "execute_activity", execute)
    monkeypatch.setattr(orchestration.workflow, "patched", lambda _: True)
    monkeypatch.setattr(orchestration.workflow, "wait_condition", wait)
    asyncio.run(instance.run({}))
    assert (
        instance.state
        == {
            "approved": "COMPLETED",
            "timeout": "COMPLETED",
            "rejected": "REJECTED",
            "cancelled": "CANCELLED",
        }[outcome]
    )
    assert ("transformation_publish" in calls) == (outcome in ("approved", "timeout"))


@pytest.mark.parametrize(
    "failure",
    ["bad_barrier", "node_refused", "activity_error", "cancel_before_node", "pause_resume"],
)
def test_legacy_transformation_stops_before_publication_on_failure_or_cancel(monkeypatch, failure):
    instance = orchestration.TransformationWorkflow()
    calls = []
    if failure == "cancel_before_node":
        instance.control({"id": "cancel", "command": "cancel"})
    if failure == "pause_resume":
        instance.control({"id": "pause", "command": "pause"})
        instance.control({"id": "pause", "command": "pause"})

    async def execute(name, context, **options):
        calls.append(name)
        if name == "transformation_load":
            return {
                "node_order": ["source"],
                "dependencies": {"source": ["missing"] if failure == "bad_barrier" else []},
            }
        if name == "transformation_node":
            if failure == "activity_error":
                raise ActivityError(
                    "fixture failed",
                    scheduled_event_id=1,
                    started_event_id=2,
                    identity="fixture",
                    activity_type=name,
                    activity_id="fixture",
                    retry_state=None,
                )
            return {"state": "BUDGET_REFUSED" if failure == "node_refused" else "COMPLETED"}
        return {"publication_id": "fixture-only"}

    async def wait(condition, **options):
        assert instance.state == "PAUSED"
        instance.control({"id": "resume", "command": "resume"})
        assert condition()

    monkeypatch.setattr(orchestration.workflow, "execute_activity", execute)
    monkeypatch.setattr(orchestration.workflow, "patched", lambda _: False)
    monkeypatch.setattr(orchestration.workflow, "wait_condition", wait)
    asyncio.run(instance.run({}))
    expected = (
        "COMPLETED"
        if failure == "pause_resume"
        else ("CANCELLED" if failure == "cancel_before_node" else "FAILED")
    )
    assert instance.status()["state"] == expected
    assert ("transformation_publish" in calls) == (failure == "pause_resume")


@pytest.mark.parametrize("mode", ["blocked_topology", "refused_result", "cancel_before_batch"])
def test_parallel_scheduler_never_publishes_unready_or_refused_nodes(monkeypatch, mode):
    instance = orchestration.TransformationWorkflow()
    calls = []
    if mode == "cancel_before_batch":
        instance.control({"id": "cancel", "command": "cancel"})

    async def execute(name, context, **options):
        calls.append(name)
        if name == "transformation_load":
            return {
                "node_order": ["source"],
                "dependencies": {"source": ["missing"] if mode == "blocked_topology" else []},
                "execution_policy": {"max_concurrent_nodes": 1},
            }
        raise AssertionError("Incomplete batch must not publish")

    async def refused():
        return {"state": "BUDGET_REFUSED", "reason": "fixture result budget"}

    monkeypatch.setattr(orchestration.workflow, "execute_activity", execute)
    monkeypatch.setattr(orchestration.workflow, "start_activity", lambda *args, **kwargs: refused())
    monkeypatch.setattr(orchestration.workflow, "patched", lambda _: True)
    asyncio.run(instance.run({}))
    assert instance.state == ("CANCELLED" if mode == "cancel_before_batch" else "FAILED")
    assert calls == ["transformation_load"]
    assert "publication" not in instance.result


@pytest.mark.parametrize("decision", ["REJECTED", "CANCELLED"])
def test_binding_review_terminal_refusal_prevents_publication(monkeypatch, decision):
    instance = orchestration.TransformationWorkflow()
    calls = []

    async def execute(name, context, **options):
        calls.append(name)
        if name == "transformation_load":
            return {"node_order": [], "dependencies": {}, "binding_review": True}
        if name == "transformation_binding_review":
            return {"state": decision}
        raise AssertionError("Rejected binding cannot reach publication")

    monkeypatch.setattr(orchestration.workflow, "execute_activity", execute)
    monkeypatch.setattr(orchestration.workflow, "patched", lambda _: True)
    asyncio.run(instance.run({}))
    assert instance.state == decision
    assert calls == ["transformation_load", "transformation_binding_review"]
