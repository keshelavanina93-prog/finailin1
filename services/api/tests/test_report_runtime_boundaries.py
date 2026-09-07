"""Report presentation, governed records, and activity lifecycle boundaries.

All source values below are explicit test fixtures, never accepted company actuals.
Temporal transport is controlled; financial results are not mocked into authority.
"""

import asyncio
import json
import os
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from fastapi import HTTPException
from fastapi.security import HTTPAuthorizationCredentials
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from temporalio.exceptions import ActivityError

from finai_api import report_workflow, workflow_worker
from finai_api.config import get_settings
from finai_api.main import app
from finai_api.security import authenticated_principal
from finai_api.services import report_workflows as records
from finai_api.services.report_export import operating_workbook
from finai_api.services.workspace import WorkspaceError


def principal():
    return authenticated_principal(
        HTTPAuthorizationCredentials(scheme="Bearer", credentials="test-token")
    )


def report_request(receipts=("fixture-source",)):
    return records.WorkflowRequest.model_validate(
        {
            "report": {
                "period": "2026-08",
                "company_label": "Synthetic workflow boundary fixture",
                "currency": "GEL",
                "receipt_ids": receipts,
            }
        }
    )


def test_export_preserves_exact_values_missingness_and_untrusted_text():
    exact = "123456789012345.123456789"
    result = {
        "state": "REFERENCE_ONLY",
        "filename": '=HYPERLINK("https://invalid.test")',
        "source_sha256": "f" * 64,
        "calculation_id": "fixture-only",
        "contract_version": "reference-fixture/1",
        "explanation": "No approved actuals",
        "missing_requirements": ["Canonical authority missing"],
        "metrics": [
            {"id": "fixture.exact", "amount": exact, "state": "REFERENCE_ONLY"},
            {"id": "fixture.zero", "amount": "0", "state": "REFERENCE_ONLY"},
            {"id": "fixture.missing", "amount": None, "state": "UNAVAILABLE"},
        ],
        "facts": [
            {
                "metric": "fixture.exact",
                "source_sheet": "=1+1",
                "source_row": 3,
                "label": '=HYPERLINK("https://invalid.test")',
                "amount": exact,
                "coordinates": ["A3", "B3"],
            }
        ],
        "comparisons": [
            {
                "metric": "fixture.exact",
                "coordinate": "B3",
                "legacy_cached": "1.00",
                "calculated": "2.00",
                "difference": "1.00",
            }
        ],
    }
    workbook = load_workbook(BytesIO(operating_workbook(result)), data_only=False)
    assert workbook.sheetnames == [
        "Operating P&L",
        "Source lineage",
        "Comparison",
        "Method & context",
    ]
    amounts = workbook["Operating P&L"]
    assert amounts["D3"].value == exact  # Presentation float never replaces exact decimal evidence.
    assert amounts["B4"].value == 0 and amounts["D4"].value == "0"
    assert amounts["B5"].value is None and amounts["D5"].value == "Unavailable"
    assert amounts["C5"].value == "UNAVAILABLE"
    assert "NOT APPROVED" in amounts["A1"].value
    assert workbook["Source lineage"]["D3"].value == result["facts"][0]["label"]
    assert workbook["Source lineage"]["F3"].value == "A3, B3"
    assert workbook["Comparison"]["E3"].value == "1.00"
    for sheet in workbook:
        assert sheet.freeze_panes == "B3"
        assert all(cell.data_type != "f" for row in sheet for cell in row)
    workbook.close()


@pytest.mark.skipif(
    os.getenv("G8_BINDING_DB_TEST") != "1" or not os.getenv("FINAI_DATABASE_URL"),
    reason="Explicit disposable PostgreSQL acceptance opt-in required",
)
def test_retained_workflow_replay_conflict_and_exact_scope():
    actor = principal()
    response = TestClient(app, headers={"Authorization": "Bearer test-token"}).post(
        "/v1/hydration/ingest",
        json={
            "scope": actor.scope.model_dump(mode="json"),
            "filename": f"workflow-fixture-{uuid4()}.csv",
            "csv_text": "account_code,debit,credit\nfixture-a,1,0\nfixture-b,0,1\n",
        },
    )
    assert response.status_code == 200, response.text
    receipt = response.json()["receipt_id"]
    request = report_request((receipt,))
    identity = records.retain(actor, request)
    assert records.retain(actor, request) == identity
    event = {"node": "coverage", "state": "FAILED", "reason": "fixture unavailable"}
    records.event(actor, identity, "coverage:0:1:failed", event)
    records.event(actor, identity, "coverage:0:1:failed", event)
    retained = records.read(actor, identity)
    assert retained["request"]["report"]["receipt_ids"] == [receipt]
    assert retained["definition"]["version"] == records.VERSION
    assert len(retained["events"]) == 1
    assert retained["events"][0]["state"] == "FAILED"
    with pytest.raises(WorkspaceError) as conflict:
        records.event(actor, identity, "coverage:0:1:failed", {**event, "state": "COMPLETED"})
    assert conflict.value.status == 409
    assert records.read(actor, identity)["events"][0]["state"] == "FAILED"
    outsider = actor.model_copy(
        update={"scope": actor.scope.model_copy(update={"period": "2026-07"})}
    )
    with pytest.raises(WorkspaceError) as hidden:
        records.read(outsider, identity)
    assert hidden.value.status == 404
    with pytest.raises(WorkspaceError) as missing_source:
        records.retain(actor, report_request(("unavailable-fixture-source",)))
    assert missing_source.value.status == 404


def test_activity_owner_is_revalidated_after_grant_change(monkeypatch):
    actor = principal()
    scope = actor.scope.model_dump(mode="json")
    assert records.current_principal(actor.actor_id, scope) == actor
    with pytest.raises(WorkspaceError) as changed_scope:
        records.current_principal(actor.actor_id, {**scope, "period": "2026-07"})
    assert changed_scope.value.status == 403
    monkeypatch.setenv(
        "FINAI_ACCESS_TOKENS",
        json.dumps(
            {
                "rotated-fixture-token": actor.model_copy(
                    update={"permissions": ("read",)}
                ).model_dump(mode="json")
            }
        ),
    )
    get_settings.cache_clear()
    assert records.current_principal(actor.actor_id, scope).permissions == ("read",)
    monkeypatch.setenv("FINAI_ACCESS_TOKENS", "{}")
    get_settings.cache_clear()
    with pytest.raises(WorkspaceError) as revoked:
        records.current_principal(actor.actor_id, scope)
    assert revoked.value.status == 403


@pytest.mark.parametrize(
    "version,expected",
    [
        ("report-source-process/1", ["report_source_coverage"]),
        ("report-source-process/2", ["report_source_hierarchy", "report_source_coverage"]),
        (
            "report-source-process/3",
            ["report_source_hierarchy", "report_source_coverage", "execution_publish"],
        ),
    ],
)
def test_workflow_version_preserves_stage_order_and_review_controls(monkeypatch, version, expected):
    instance = report_workflow.ReportSourceWorkflow()
    assert instance.status() == {"state": "QUEUED", "result": {}}
    calls, states = [], []
    commands = iter(["pause", "retry", "resume", "complete"])

    async def execute(name, context, **options):
        calls.append(name)
        assert context == {
            "definition_version": version,
            "workflow_id": "opaque-fixture",
            "generation": 0,
        }
        assert options["retry_policy"].maximum_attempts == 3
        if name == "report_source_coverage":
            assert options["retry_policy"].initial_interval.total_seconds() == 2
            return {"assessment_id": "fixture-unavailable", "state": "UNAVAILABLE"}
        return {"reference": "fixture-only"}

    async def wait(condition):
        states.append(instance.state)
        command = next(commands)
        instance.control({"id": command, "command": command})
        instance.control(
            {"id": command, "command": command}
        )  # Duplicate delivery cannot repeat it.
        assert len(instance.commands) == 1 and condition()

    monkeypatch.setattr(report_workflow.workflow, "execute_activity", execute)
    monkeypatch.setattr(report_workflow.workflow, "wait_condition", wait)
    result = asyncio.run(
        instance.run({"definition_version": version, "workflow_id": "opaque-fixture"})
    )
    assert calls == expected
    assert states == ["WAITING_REVIEW", "PAUSED", "PAUSED", "WAITING_REVIEW"]
    assert instance.status()["state"] == "REVIEWED"
    assert result["state"] == "UNAVAILABLE"  # Review control does not certify missing data.


def test_workflow_exhausted_failure_requires_retry_and_can_cancel(monkeypatch):
    instance = report_workflow.ReportSourceWorkflow()
    generations, states = [], []
    commands = iter(["complete", "retry", "cancel"])

    async def execute(name, context, **options):
        generations.append(context["generation"])
        if context["generation"] == 0:
            raise ActivityError(
                "fixture exhausted",
                scheduled_event_id=1,
                started_event_id=2,
                identity="fixture",
                activity_type=name,
                activity_id="fixture",
                retry_state=None,
            )
        return {"assessment_id": "fixture-only", "state": "UNAVAILABLE"}

    async def wait(condition):
        states.append(instance.state)
        command = next(commands)
        instance.control({"id": command, "command": command})
        assert condition()

    monkeypatch.setattr(report_workflow.workflow, "execute_activity", execute)
    monkeypatch.setattr(report_workflow.workflow, "wait_condition", wait)
    assert asyncio.run(instance.run({}))["state"] == "UNAVAILABLE"
    assert generations == [0, 1]
    assert states == ["FAILED", "FAILED", "WAITING_REVIEW"]
    assert instance.state == "CANCELLED"


@pytest.fixture
def worker_boundary(monkeypatch):
    actor = principal()
    context = {
        "actor_id": actor.actor_id,
        "scope": actor.scope.model_dump(mode="json"),
        "workflow_id": "fixture-only",
        "generation": 2,
    }
    events, stages = [], []
    record = {
        "request": report_request(("b", "a", "a")).model_dump(mode="json"),
        "definition": records.DEFINITION,
    }
    monkeypatch.setattr(records, "read", lambda _actor, _identity: record)
    monkeypatch.setattr(records, "event", lambda *args: events.append(args))
    monkeypatch.setattr(workflow_worker.publication, "stage", lambda *args: stages.append(args))
    monkeypatch.setattr(workflow_worker.activity, "info", lambda: SimpleNamespace(attempt=3))
    return SimpleNamespace(
        actor=actor, context=context, events=events, stages=stages, record=record
    )


def test_hierarchy_keeps_source_proofs_out_of_temporal_history(monkeypatch, worker_boundary):
    boundary = worker_boundary
    seen = []

    def source(_actor, receipt):
        seen.append(receipt)
        return SimpleNamespace(
            receipt=SimpleNamespace(
                classifier_version="1c-biff-tb-layout/1"
                if receipt == "a"
                else "unrelated-layout/1",
                source_sha256="a" * 64,
                candidates=[],
            )
        )

    monkeypatch.setattr(workflow_worker, "detail", source)
    output = workflow_worker.hierarchy(boundary.context)
    assert seen == ["a", "b"]
    assert output == {"event_id": "hierarchy:2:3:completed", "sources_checked": 1}
    proof = boundary.stages[0][-1]["proofs"][0]["proof"]
    assert proof["state"] == "REVIEW_REQUIRED" and proof["selected_rows"] == []
    assert "proofs" not in output
    assert [event[-1]["state"] for event in boundary.events] == ["RUNNING", "COMPLETED"]


@pytest.mark.parametrize("activity_name", ["hierarchy", "coverage"])
def test_activity_failure_records_attempt_without_exception_payload_leak(
    monkeypatch, worker_boundary, activity_name
):
    boundary = worker_boundary

    def fail(*args):
        raise RuntimeError("fixture-private-source-payload")

    monkeypatch.setattr(workflow_worker, "detail", fail)
    monkeypatch.setattr(workflow_worker, "retain_assessment", fail)
    with pytest.raises(RuntimeError, match="fixture-private-source-payload"):
        getattr(workflow_worker, activity_name)(boundary.context)
    assert [event[-1]["state"] for event in boundary.events] == ["RUNNING", "FAILED"]
    assert "fixture-private-source-payload" not in json.dumps(
        [event[-1] for event in boundary.events]
    )
    assert not boundary.stages


def test_coverage_stages_only_assessment_reference_and_legacy_does_not_publish(
    monkeypatch, worker_boundary
):
    boundary = worker_boundary
    monkeypatch.setattr(
        workflow_worker,
        "retain_assessment",
        lambda _actor, request: {
            "assessment_id": "fixture-assessment",
            "state": "UNAVAILABLE",
            "private_source_payload": "never-in-temporal",
        },
    )
    expected = {"assessment_id": "fixture-assessment", "state": "UNAVAILABLE"}
    assert workflow_worker.coverage(boundary.context) == expected
    assert boundary.stages[0][-1] == expected
    boundary.record["definition"] = {"version": "report-source-process/1"}
    boundary.stages.clear()
    assert workflow_worker.coverage(boundary.context) == expected
    assert not boundary.stages
    monkeypatch.setattr(
        workflow_worker.publication,
        "publish",
        lambda *args: {
            "publication_id": "fixture-publication",
            "generation": 2,
            "private_source_payload": "hidden",
        },
    )
    assert workflow_worker.publish_outputs(boundary.context) == {
        "publication_id": "fixture-publication",
        "generation": 2,
    }


@pytest.mark.parametrize("activity_name", ["hierarchy", "coverage"])
def test_activity_permission_revocation_prevents_work(monkeypatch, worker_boundary, activity_name):
    boundary = worker_boundary
    monkeypatch.setenv(
        "FINAI_ACCESS_TOKENS",
        json.dumps(
            {
                "test-token": boundary.actor.model_copy(
                    update={"permissions": ("read",)}
                ).model_dump(mode="json")
            }
        ),
    )
    get_settings.cache_clear()
    with pytest.raises(HTTPException) as denied:
        getattr(workflow_worker, activity_name)(boundary.context)
    assert denied.value.status_code == 403
    assert not boundary.events and not boundary.stages

