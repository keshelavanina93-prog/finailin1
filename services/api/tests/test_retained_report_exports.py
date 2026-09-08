"""Synthetic retained values only: exports neither query nor calculate financial truth."""

import base64
import hashlib
import json
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from finai_api.services.retained_report_exports import build, validate_artifact_bytes

PIN = {
    "resource_id": "00000000-0000-4000-8000-000000000001",
    "version_id": "00000000-0000-4000-8000-000000000002",
    "content_hash": "a" * 64,
}


def snapshot(row_count=8, section_count=1):
    fields = [
        {
            "key": "amount",
            "label": "ზუსტი თანხა",
            "kind": "decimal",
            "unit": "GEL",
            "definition": PIN,
        },
        {"key": "name", "label": "名称", "kind": "text", "definition": PIN},
        {"key": "hidden", "label": "PRIVATE_COLUMN", "kind": "text", "definition": PIN},
    ]
    values = [
        {"state": "VALUE", "value": "12345678901234567890.12345678901234567890"},
        {"state": "VALUE", "value": "0.000000000000000000000001"},
        {"state": "MISSING"},
        {"state": "NULL", "value": None},
        {"state": "VALUE", "value": ""},
        {"state": "VALUE", "value": False},
        {"state": "VALUE", "value": 0},
        {"state": "VALUE", "value": None},
    ]
    sections = []
    for section_index in range(section_count):
        reference = {
            "section_id": str(section_index),
            "title": f"Section title {section_index}",
            "invocation_id": "synthetic-invocation",
            "receipt_hash": "b" * 64,
            "descriptor_sha256": "c" * 64,
            "columns": ["name", "amount"],
            "filters": [],
            "group_by": None,
        }
        rows = [
            {
                "key": f"row_{index:064x}",
                "label": f"UNSELECTED_ROW_LABEL_{index}",
                "values": {
                    "amount": deepcopy(values[index % len(values)]),
                    "name": {"state": "VALUE", "value": f"s{section_index}-r{index}"},
                    "hidden": {"state": "VALUE", "value": "PRIVATE_VALUE"},
                },
                "contributor_count": 1,
                "trace": PIN,
            }
            for index in range(row_count)
        ]
        descriptor = {
            "fields": fields,
            "context": [
                {"label": "Period", "value": f"Period {section_index}"},
                {"label": "Book", "value": f"Book {section_index}"},
                {"label": "Currency", "value": "GEL"},
            ],
            "coverage": [{"label": "Retained rows", "value": str(row_count)}],
            "valid_at": f"2026-01-0{section_index + 1}T00:00:00Z",
            "known_at": f"2026-02-0{section_index + 1}T00:00:00Z",
            "recorded_at": f"2026-03-0{section_index + 1}T00:00:00Z",
            "authority": "Retained only",
            "function": PIN,
            "company": PIN,
            "definitions": [PIN],
            "run_id": "synthetic-run",
            "unavailable_operations": [],
            "current_use_authorized": False,
            "business_effect_authorized": False,
        }
        sections.append(
            {
                "reference": reference,
                "projection": {"descriptor": descriptor, "rows": rows, "total_rows": row_count},
                "contributors": {
                    row["key"]: [
                        {
                            "reference": PIN,
                            "source_sha256": "d" * 64,
                            "sheet": "原本",
                            "coordinate": "A1",
                            "cells": [{"value": "PRIVATE_EVIDENCE_CELL"}],
                        }
                    ]
                    for row in rows
                },
                "authority_observation": {"current_use_authorized": False},
                "current_use_authorized": False,
                "business_effect_authorized": False,
            }
        )
    return {
        "contract": "retained-report-snapshot/1",
        "company": PIN,
        "company_label": "სინთეზური კომპანია / 測試",
        "composition": {
            "company_id": PIN["resource_id"],
            "valid_at": "2026-04-01T00:00:00Z",
            "known_at": "2026-05-01T00:00:00Z",
            "title": "Synthetic retained report",
            "commentary": "Author commentary",
            "sections": [s["reference"] for s in sections],
        },
        "sections": sections,
        "current_use_authorized": False,
        "business_effect_authorized": False,
    }


def workbook(artifacts):
    return load_workbook(BytesIO(validate_artifact_bytes(artifacts["xlsx"])), data_only=False)


def grid(sheet):
    _, header, last_column, last_row = range_boundaries(sheet.auto_filter.ref)
    return list(
        sheet.iter_rows(min_row=header, max_row=last_row, max_col=last_column, values_only=True)
    )


def records(sheet):
    grouped = {}
    for section, row_key, record, part, text in sheet.iter_rows(min_row=2, values_only=True):
        key = (section, row_key, record)
        grouped.setdefault(key, []).append((int(part), text))
    return {
        key: json.loads("".join(text for _, text in sorted(parts)))
        for key, parts in grouped.items()
    }


def test_repeat_is_byte_identical_without_mutation_or_temp_files(monkeypatch):
    def prohibited(*args, **kwargs):
        raise AssertionError("Exporter attempted worksheet disk staging")

    monkeypatch.setattr("openpyxl.worksheet._writer.create_temporary_file", prohibited)
    original = snapshot()
    before = deepcopy(original)
    first, second = build(original), build(original)
    assert original == before
    assert first == second
    for artifact in first.values():
        content = validate_artifact_bytes(artifact)
        assert artifact["size_bytes"] == len(content)
        assert artifact["sha256"] == hashlib.sha256(content).hexdigest()
    with ZipFile(BytesIO(validate_artifact_bytes(first["xlsx"]))) as archive:
        assert archive.namelist() == sorted(archive.namelist())
        assert {item.date_time for item in archive.infolist()} == {(2000, 1, 1, 0, 0, 0)}
        core = archive.read("docProps/core.xml")
        assert core.count(b"2000-01-01T00:00:00Z") == 2


def test_all_selected_rows_columns_order_and_independent_context():
    artifacts = build(snapshot(row_count=600, section_count=2))
    book = workbook(artifacts)
    assert book.sheetnames == ["Report", "Section 1", "Section 2", "Advanced"]
    document = validate_artifact_bytes(artifacts["html"]).decode("utf-8")
    for index in range(2):
        sheet = book[f"Section {index + 1}"]
        data = grid(sheet)
        assert data[0] == ("名称", "ზუსტი თანხა (GEL)")
        assert len(data) == 601
        assert [row[0] for row in data[1:]] == [f"s{index}-r{row}" for row in range(600)]
        for row in range(600):
            assert f"<td>s{index}-r{row}</td>" in document
        for expected in (f"Period {index}", f"Book {index}", f"2026-03-0{index + 1}T00:00:00Z"):
            assert expected in document
            assert any(expected == value for row in sheet.values for value in row)
    for secret in (
        "PRIVATE_COLUMN",
        "PRIVATE_VALUE",
        "PRIVATE_EVIDENCE_CELL",
        "UNSELECTED_ROW_LABEL",
    ):
        assert secret not in document
        assert not any(
            secret in str(value) for sheet in book for row in sheet.values for value in row
        )


def test_exact_precision_and_distinct_states():
    original = snapshot()
    artifacts = build(original)
    book = workbook(artifacts)
    values = [row[1] for row in grid(book["Section 1"])[1:]]
    assert values == [
        "12345678901234567890.12345678901234567890",
        "0.000000000000000000000001",
        "[Missing]",
        "[Null]",
        "[Empty text]",
        "False",
        "0",
        "[Null value]",
    ]
    exact = records(book["Advanced"])
    for row in original["sections"][0]["projection"]["rows"]:
        cell = exact[("1", row["key"], "amount")]
        assert cell["cell"] == row["values"]["amount"]
        assert cell["value_key_present"] == ("value" in row["values"]["amount"])
    assert exact[("1", "row_" + f"{5:064x}", "amount")]["value_type"] == "bool"
    assert exact[("1", "row_" + f"{6:064x}", "amount")]["value_type"] == "int"


def test_formula_strings_and_html_are_inert():
    original = snapshot()
    attacks = [
        '=HYPERLINK("https://example.test","go")',
        "+1+1",
        "-1+1",
        "@SUM(A1:A2)",
        '<script>alert("x")</script>',
        '<img src="https://example.test/x" onerror="alert(1)">',
        "ქართული / 中文",
        "https://example.test/external.xlsx",
    ]
    original["composition"]["title"] = attacks[0]
    original["composition"]["commentary"] = attacks[4]
    for row, attack in zip(original["sections"][0]["projection"]["rows"], attacks, strict=True):
        row["values"]["name"]["value"] = attack
    artifacts = build(original)
    book = workbook(artifacts)
    assert [row[0] for row in grid(book["Section 1"])[1:]] == attacks
    for sheet in book:
        for row in sheet:
            for cell in row:
                assert cell.data_type != "f"
                assert cell.hyperlink is None
    with ZipFile(BytesIO(validate_artifact_bytes(artifacts["xlsx"]))) as archive:
        assert not any("externalLink" in name for name in archive.namelist())
        assert not any(
            b"<f>" in archive.read(name) for name in archive.namelist() if name.endswith(".xml")
        )
    document = validate_artifact_bytes(artifacts["html"]).decode("utf-8")
    assert "<script>" not in document and "<img " not in document
    assert "&lt;script&gt;" in document
    assert "Content-Security-Policy" in document and "thead{display:table-header-group}" in document


def test_absent_cells_context_and_oversized_text_remain_explicit():
    original = snapshot(row_count=1)
    section = original["sections"][0]
    section["projection"]["descriptor"]["context"] = []
    row = section["projection"]["rows"][0]
    del row["values"]["amount"]
    row["values"]["name"]["value"] = "文" * 34000 + "\x0b"
    artifacts = build(original)
    book = workbook(artifacts)
    assert grid(book["Section 1"])[1][1] == "[Missing]"
    assert grid(book["Section 1"])[1][0].endswith("[continued in Advanced]")
    exact = records(book["Advanced"])
    assert exact[("1", row["key"], "name")]["cell"] == row["values"]["name"]
    assert exact[("1", row["key"], "amount")]["selected_cell_present"] is False
    document = validate_artifact_bytes(artifacts["html"]).decode("utf-8")
    assert "<dt>Period</dt><dd>Not supplied</dd>" in document
    assert "<dt>Book</dt><dd>Not supplied</dd>" in document
    assert "<dt>Currency</dt><dd>Not supplied</dd>" in document


@pytest.mark.parametrize(
    "change",
    [
        {"content_base64": "!invalid!"},
        {"size_bytes": 123},
        {"sha256": "0" * 64},
        {"content_base64": base64.b64encode(b"different").decode("ascii")},
    ],
)
def test_corrupt_artifact_is_rejected(change):
    artifact = build(snapshot(row_count=1))["html"]
    artifact.update(change)
    with pytest.raises(ValueError, match="Invalid retained report artifact"):
        validate_artifact_bytes(artifact)
